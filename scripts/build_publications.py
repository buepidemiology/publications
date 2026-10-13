"""
Build the group publication pages from public ORCID records ONLY.

The only service contacted is the ORCID public API (pub.orcid.org).
No PubMed, Crossref, Scopus or any other source is queried.

Roster:  ORCIDs.xlsx (top level of the repo, or in data/), one sheet per group ("Faculty", "PhD", "Post Docs").
         Column A = Name ("Last, First"), column B = ORCID iD.
         People without an ORCID iD are skipped (listed in the Action log).

Output:  docs/faculty.html, docs/phd.html, docs/postdocs.html  (the three pages)
         docs/index.html          (redirects to the Faculty page)
         docs/publications.json   (all data)
"""

import html
import json
import re
import sys
import time
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
# The roster may live in data/ORCIDs.xlsx or at the top level of the repo.
ROSTER_CANDIDATES = [ROOT / "data" / "ORCIDs.xlsx", ROOT / "ORCIDs.xlsx"]
ROSTER_PATH = next((p for p in ROSTER_CANDIDATES if p.exists()), ROSTER_CANDIDATES[0])
OUTPUT_DIR = ROOT / "docs"
JSON_OUTPUT = OUTPUT_DIR / "publications.json"

ORCID_API = "https://pub.orcid.org/v3.0"
REQUEST_TIMEOUT = 30
REQUEST_PAUSE_SECONDS = 0.3
BULK_SIZE = 100            # ORCID allows up to 100 works per bulk request
MAX_AUTHORS_SHOWN = 10     # longer author lists end with "et al."
NEWEST_COUNT = 10          # how many works appear in "Newest publications" at the top
NEW_BADGE_DAYS = 7         # works first seen within this many days get a "New" badge

SITE_TITLE = "BU Epidemiology Group Publications"

# Order here = order of the tabs on every page.
GROUPS = [
    {"key": "faculty",  "sheet": "faculty",  "title": "Faculty",      "file": "faculty.html"},
    {"key": "phd",      "sheet": "phd",      "title": "PhD Students", "file": "phd.html"},
    {"key": "postdocs", "sheet": "postdocs", "title": "Postdocs",     "file": "postdocs.html"},
]

ORCID_PATTERN = re.compile(r"(\d{4}-\d{4}-\d{4}-\d{3}[\dXx])")

SESSION = requests.Session()
SESSION.headers.update({
    "Accept": "application/json",
    "User-Agent": "BU-Epi-group-publications (GitHub Actions)",
})


# ------------------------------------------------------------------ roster

def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def sheet_key(name):
    return re.sub(r"[^a-z0-9]", "", name.lower())


def orcid_checksum_ok(orcid):
    digits = orcid.replace("-", "")
    total = 0
    for ch in digits[:-1]:
        total = (total + int(ch)) * 2
    result = (12 - total % 11) % 11
    return digits[-1].upper() == ("X" if result == 10 else str(result))


def display_name(name):
    """'Sixpence, Alick' -> 'Alick Sixpence'."""
    if "," in name:
        last, first = [part.strip() for part in name.split(",", 1)]
        return f"{first} {last}".strip()
    return name


def slugify(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def read_roster():
    if not ROSTER_PATH.exists():
        raise RuntimeError("Roster not found. Put ORCIDs.xlsx at the top of the repo or in data/.")
    workbook = load_workbook(ROSTER_PATH, read_only=True, data_only=True)
    sheets = {sheet_key(name): workbook[name] for name in workbook.sheetnames}
    roster = {}

    for group in GROUPS:
        sheet = sheets.get(group["sheet"])
        if sheet is None:
            raise RuntimeError(
                f"No sheet for '{group['title']}' in {ROSTER_PATH.name}. "
                f"Sheets found: {workbook.sheetnames}"
            )

        people, skipped, seen = [], [], set()
        for row in sheet.iter_rows(values_only=True):
            name = clean(row[0] if len(row) > 0 else "")
            raw_orcid = clean(row[1] if len(row) > 1 else "")
            if not name or name.lower() == "name":
                continue

            match = ORCID_PATTERN.search(raw_orcid)
            if not match:
                skipped.append(f"{name}: no ORCID iD")
                continue
            orcid = match.group(1).upper()
            if not orcid_checksum_ok(orcid):
                skipped.append(f"{name}: ORCID iD {orcid} fails checksum (typo?)")
                continue
            if orcid in seen:
                continue
            seen.add(orcid)
            people.append({"name": display_name(name), "orcid": orcid, "sort": name.lower()})

        people.sort(key=lambda person: person.pop("sort"))
        roster[group["key"]] = {"people": people, "skipped": skipped}

    return roster


# ------------------------------------------------------------------- ORCID

def api_get(path):
    last_error = None
    for attempt in range(4):
        try:
            response = SESSION.get(f"{ORCID_API}{path}", timeout=REQUEST_TIMEOUT)
            if response.status_code == 404:
                return None
            if response.status_code in (429, 500, 502, 503, 504):
                last_error = f"HTTP {response.status_code}"
            else:
                response.raise_for_status()
                return response.json()
        except requests.RequestException as error:
            last_error = str(error)
        time.sleep(2 * (attempt + 1))
    raise RuntimeError(last_error or "ORCID request failed")


def value_at(data, *keys):
    for key in keys:
        if not isinstance(data, dict):
            return ""
        data = data.get(key)
    return data.strip() if isinstance(data, str) else ""


def clean_text(text):
    """Remove markup like <scp> or <i> and decode entities such as &amp;amp;."""
    text = text or ""
    for _ in range(2):
        text = html.unescape(text)
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", text)).strip()


def normalize_doi(value):
    doi = (value or "").strip().lower()
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi)
    doi = re.sub(r"^doi:\s*", "", doi)
    return doi.rstrip(" .;,")


def external_ids(*records):
    found = {}
    for record in records:
        for ext in value_list(record, "external-ids", "external-id"):
            if (ext.get("external-id-relationship") or "self") != "self":
                continue
            kind = (ext.get("external-id-type") or "").lower()
            value = value_at(ext, "external-id-normalized", "value") or value_at(ext, "external-id-value")
            if kind and value and kind not in found:
                found[kind] = {"value": value, "url": value_at(ext, "external-id-url", "value")}
    return found


def value_list(data, *keys):
    for key in keys:
        if not isinstance(data, dict):
            return []
        data = data.get(key)
    return [item for item in (data or []) if isinstance(item, dict)]


def make_publication(summary, group):
    title = clean_text(value_at(summary, "title", "title", "value"))
    if not title:
        return None

    ids = external_ids(summary, group)
    doi = normalize_doi((ids.get("doi") or {}).get("value"))
    year = value_at(summary, "publication-date", "year", "value")
    month = value_at(summary, "publication-date", "month", "value")
    day = value_at(summary, "publication-date", "day", "value")

    url = value_at(summary, "url", "value")
    if not url:
        url = next((info["url"] for info in ids.values() if info.get("url")), "")

    return {
        "put_code": str(summary.get("put-code") or ""),
        "title": title,
        "year": int(year) if year.isdigit() else 0,
        "month": int(month) if month.isdigit() else 0,
        "day": int(day) if day.isdigit() else 0,
        "doi": doi,
        "journal": clean_text(value_at(summary, "journal-title", "value")),
        "type": summary.get("type") or "",
        "authors": [],
        "url": url,
    }


def author_names(work):
    names, seen = [], set()
    for contributor in value_list(work, "contributors", "contributor"):
        name = clean_text(value_at(contributor, "credit-name", "value"))
        key = re.sub(r"[^a-z0-9]", "", name.lower())
        if name and key not in seen:  # ORCID records often repeat the owner's name
            seen.add(key)
            names.append(name)
    return names


def add_full_details(orcid, publications):
    """Fetch full work records in bulk (100 per request) to get author lists."""
    by_code = {pub["put_code"]: pub for pub in publications if pub["put_code"]}
    codes = list(by_code)

    for start in range(0, len(codes), BULK_SIZE):
        chunk = codes[start:start + BULK_SIZE]
        try:
            data = api_get(f"/{orcid}/works/{','.join(chunk)}") or {}
        except Exception as error:
            print(f"    author lists unavailable for {len(chunk)} works: {error}")
            continue

        for item in data.get("bulk") or []:
            work = (item or {}).get("work")
            if not work:
                continue
            pub = by_code.get(str(work.get("put-code")))
            if not pub:
                continue
            pub["authors"] = author_names(work)
            subtitle = clean_text(value_at(work, "title", "subtitle", "value"))
            if subtitle and subtitle.lower() not in pub["title"].lower():
                pub["title"] = f"{pub['title']}: {subtitle}"
            if not pub["journal"]:
                pub["journal"] = clean_text(value_at(work, "journal-title", "value"))

        time.sleep(REQUEST_PAUSE_SECONDS)


def publication_key(pub):
    if pub["doi"]:
        return f"doi:{pub['doi']}"
    title = re.sub(r"[^a-z0-9]", "", pub["title"].lower())
    return f"title:{title}|year:{pub['year']}"


def publication_date(pub):
    return (pub.get("year") or 0, pub.get("month") or 0, pub.get("day") or 0)


def completeness(pub):
    return (len(pub["authors"]), bool(pub["journal"]), bool(pub["year"]))


def get_orcid_works(orcid):
    data = api_get(f"/{orcid}/works")
    if not data:
        return []

    publications = []
    for group in data.get("group") or []:
        summaries = value_list(group, "work-summary")
        if not summaries:
            continue
        # ORCID groups duplicate versions of one work; the first is the owner's preferred one.
        pub = make_publication(summaries[0], group)
        if pub:
            publications.append(pub)

    add_full_details(orcid, publications)

    unique = {}
    for pub in publications:
        key = publication_key(pub)
        if key not in unique or completeness(pub) > completeness(unique[key]):
            unique[key] = pub

    return sorted(unique.values(), key=lambda p: (publication_date(p), p["title"].lower()), reverse=True)


# ------------------------------------------------------------------- HTML

def format_type(work_type):
    return (work_type or "other").replace("-", " ").replace("_", " ").capitalize()


def end_sentence(text):
    text = text.strip()
    return text if text.endswith((".", "?", "!")) else text + "."


def citation_html(pub):
    authors = pub["authors"]
    parts = []
    if authors:
        shown = ", ".join(authors[:MAX_AUTHORS_SHOWN])
        if len(authors) > MAX_AUTHORS_SHOWN:
            shown += ", et al"
        parts.append(html.escape(end_sentence(shown)))

    parts.append(f'<span class="title">{html.escape(end_sentence(pub["title"]))}</span>')

    venue = html.escape(pub["journal"])
    if venue:
        venue = f"<em>{venue}</em>"
    if pub["year"]:
        month = MONTHS[pub["month"] - 1] + " " if 1 <= (pub.get("month") or 0) <= 12 else ""
        venue = f"{venue} ({month}{pub['year']})".strip()
    if venue:
        parts.append(venue + ".")

    return " ".join(parts)


MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def is_new(pub):
    try:
        first_seen = date.fromisoformat(pub.get("first_seen") or "")
    except ValueError:
        return False
    return (datetime.now(timezone.utc).date() - first_seen).days < NEW_BADGE_DAYS


def render_publication(pub, members=None):
    meta = []
    if is_new(pub):
        meta.append('<span class="new">New</span>')
    if members:
        meta.append(f'<span class="members">{html.escape(", ".join(members))}</span>')
    meta.append(f'<span class="type">{html.escape(format_type(pub["type"]))}</span>')
    if pub["doi"]:
        doi = html.escape(pub["doi"], quote=True)
        meta.append(f'<a href="https://doi.org/{doi}" target="_blank" rel="noopener noreferrer">DOI</a>')
    elif pub["url"]:
        url = html.escape(pub["url"], quote=True)
        meta.append(f'<a href="{url}" target="_blank" rel="noopener noreferrer">View record</a>')

    return (f'        <li>{citation_html(pub)}'
            f'<div class="publication-meta">{" · ".join(meta)}</div></li>')


def render_researcher(person):
    name = html.escape(person["name"])
    orcid = html.escape(person["orcid"])
    works = person["works"]
    count = len(works)

    lines = [
        f'  <section class="researcher" id="{person["slug"]}">',
        f'    <h2>{name} <span class="count">({count} work{"" if count == 1 else "s"})</span></h2>',
        f'    <p class="orcid">ORCID: <a href="https://orcid.org/{orcid}" target="_blank" '
        f'rel="noopener noreferrer">{orcid}</a></p>',
    ]
    if person["status"] == "stale":
        lines.append('    <p class="none">ORCID could not be reached on the last update; '
                     "showing the most recently saved list.</p>")

    if not works:
        lines.append('    <p class="none">No public works are listed on this ORCID record yet.</p>')
    else:
        by_year = defaultdict(list)
        for pub in works:
            by_year[pub["year"]].append(pub)
        for year in sorted(by_year, reverse=True):
            label = str(year) if year else "Year not listed"
            lines.append(f'    <div class="year"><h3>{label}</h3>')
            lines.append("      <ol>")
            lines.extend(render_publication(pub) for pub in by_year[year])
            lines.append("      </ol></div>")

    lines.append("  </section>")
    return lines


def render_newest(researchers):
    """The group's most recent works by publication date, each listed once."""
    merged = {}
    for person in researchers:
        for pub in person["works"]:
            if not pub["year"]:
                continue
            key = publication_key(pub)
            entry = merged.setdefault(key, {"pub": pub, "members": []})
            if person["name"] not in entry["members"]:
                entry["members"].append(person["name"])

    newest = sorted(merged.values(),
                    key=lambda e: (publication_date(e["pub"]), e["pub"]["title"].lower()),
                    reverse=True)[:NEWEST_COUNT]
    if not newest:
        return []

    lines = ['  <section id="newest">',
             "    <h2>Newest publications</h2>",
             "    <ol>"]
    lines.extend(render_publication(e["pub"], e["members"]) for e in newest)
    lines.extend(["    </ol>", "  </section>"])
    return lines


STYLE = """
    :root { --text: #1f2937; --muted: #4b5563; --rule: #d1d5db; --accent: #0f4c81;
            --link: #0757c8; --bg: #ffffff; --chip: #edf7ed; --chip-text: #256029;
            --highlight: #eef4fb; }
    @media (prefers-color-scheme: dark) {
      :root { --text: #e5e7eb; --muted: #9ca3af; --rule: #374151; --accent: #6ea8dc;
              --link: #8ab4f8; --bg: #111827; --chip: #1f3b25; --chip-text: #a7d7ae;
              --highlight: #17233a; }
    }
    body { font-family: Arial, Helvetica, sans-serif; color: var(--text); background: var(--bg);
           line-height: 1.6; max-width: 1080px; margin: 2rem auto; padding: 0 1rem 4rem; }
    a { color: var(--link); }
    h1 { margin-bottom: 0.2rem; }
    nav.tabs { display: flex; gap: 0.25rem; border-bottom: 2px solid var(--rule); margin: 1.2rem 0; flex-wrap: wrap; }
    nav.tabs a { padding: 0.5rem 1rem; text-decoration: none; color: var(--muted); font-weight: bold;
                 border-bottom: 3px solid transparent; margin-bottom: -2px; }
    nav.tabs a.active { color: var(--accent); border-bottom-color: var(--accent); }
    .meta, .orcid, .publication-meta, .none { color: var(--muted); font-size: 0.92rem; }
    #search { width: 100%; max-width: 420px; padding: 0.5rem 0.7rem; font: inherit; color: var(--text);
              background: var(--bg); border: 1px solid var(--rule); border-radius: 6px; }
    .jump { columns: 3 14rem; padding-left: 1.2rem; font-size: 0.95rem; }
    h2 { margin: 3rem 0 0.3rem; padding-top: 0.7rem; border-top: 2px solid var(--accent); }
    h3 { margin: 1.25rem 0 0.35rem; border-bottom: 1px solid var(--rule); padding-bottom: 0.25rem; font-size: 1.05rem; }
    ol { padding-left: 1.5rem; }
    li { margin-bottom: 1rem; }
    .title { font-weight: bold; }
    .publication-meta { margin-top: 0.2rem; }
    .type { display: inline-block; border-radius: 999px; padding: 0.05rem 0.5rem;
            background: var(--chip); color: var(--chip-text); }
    .count { color: var(--muted); font-weight: normal; font-size: 1rem; }
    .hidden { display: none; }
    #newest { margin: 1.5rem 0 1rem; padding: 0.5rem 1.25rem 0.25rem; border-radius: 10px;
              background: var(--highlight); border-left: 5px solid var(--accent); }
    #newest h2 { margin: 0.6rem 0 0.8rem; padding: 0; border: 0; }
    .new { display: inline-block; border-radius: 999px; padding: 0.05rem 0.5rem; font-weight: bold;
           background: var(--accent); color: var(--bg); }
    .members { font-weight: bold; color: var(--text); }
"""

SEARCH_SCRIPT = """
  <script>
    (function () {
      var box = document.getElementById('search');
      box.addEventListener('input', function () {
        var q = box.value.trim().toLowerCase();
        var newest = document.getElementById('newest');
        if (newest) { newest.classList.toggle('hidden', !!q); }
        document.querySelectorAll('section.researcher').forEach(function (section) {
          var nameHit = section.querySelector('h2').textContent.toLowerCase().indexOf(q) !== -1;
          var anyVisible = false;
          section.querySelectorAll('.year').forEach(function (year) {
            var yearVisible = false;
            year.querySelectorAll('li').forEach(function (li) {
              var show = !q || nameHit || li.textContent.toLowerCase().indexOf(q) !== -1;
              li.classList.toggle('hidden', !show);
              if (show) { yearVisible = true; }
            });
            year.classList.toggle('hidden', !yearVisible);
            if (yearVisible) { anyVisible = true; }
          });
          section.classList.toggle('hidden', !!q && !nameHit && !anyVisible);
        });
      });
    })();
  </script>"""


def render_page(group_key, groups, generated_at):
    group = groups[group_key]
    researchers = group["researchers"]
    total = sum(len(person["works"]) for person in researchers)

    active = ' class="active"'
    tabs = "".join(
        f'<a href="{groups[key]["file"]}"{active if key == group_key else ""}>'
        f'{html.escape(groups[key]["title"])}</a>'
        for key in groups
    )
    jump = "".join(
        f'<li><a href="#{person["slug"]}">{html.escape(person["name"])}</a> '
        f'<span class="count">({len(person["works"])})</span></li>'
        for person in researchers
    )

    lines = [
        "<!doctype html>",
        '<html lang="en">',
        "<head>",
        '  <meta charset="utf-8">',
        '  <meta name="viewport" content="width=device-width, initial-scale=1">',
        f"  <title>{html.escape(group['title'])} | {SITE_TITLE}</title>",
        f"  <style>{STYLE}  </style>",
        "</head>",
        "<body>",
        f"  <h1>{SITE_TITLE}</h1>",
        f'  <p class="meta">Public works listed on each person\'s ORCID record. '
        f"Last updated: {html.escape(generated_at)}.</p>",
        f'  <nav class="tabs">{tabs}</nav>',
        f'  <p class="meta">{len(researchers)} {html.escape(group["title"].lower())} · {total} works</p>',
        '  <input id="search" type="search" placeholder="Search names, titles, journals…" '
        'aria-label="Search publications">',
    ]
    lines.extend(render_newest(researchers))
    lines.append(f'  <h2 style="border:0">Publications by person</h2>')
    lines.append(f'  <ul class="jump">{jump}</ul>')
    for person in researchers:
        lines.extend(render_researcher(person))
    lines.append('  <p class="meta" style="margin-top:3rem">Missing a publication? '
                 "Add it to your ORCID record; it will appear after the next weekly update.</p>")
    lines.append(SEARCH_SCRIPT)
    lines.extend(["</body>", "</html>"])
    return "\n".join(lines)


def render_redirect(target):
    return (f'<!doctype html>\n<html lang="en"><head><meta charset="utf-8">'
            f'<meta http-equiv="refresh" content="0; url={target}">'
            f"<title>{SITE_TITLE}</title></head>"
            f'<body><p><a href="{target}">Go to group publications</a></p></body></html>\n')


# -------------------------------------------------------------------- main

def load_previous():
    """Last saved works per ORCID iD (fallback if ORCID is unreachable) and the
    date each work was first seen (for the "New" badge)."""
    try:
        previous = json.loads(JSON_OUTPUT.read_text(encoding="utf-8"))
    except Exception:
        previous = {}
    saved, first_seen = {}, {}
    groups = previous.get("groups") if isinstance(previous.get("groups"), dict) else {}
    for group in groups.values():
        for person in group.get("researchers") or []:
            orcid = person.get("orcid")
            if not orcid:
                continue
            if person.get("works"):
                saved[orcid] = person["works"]
            for pub in person.get("works") or []:
                first_seen[(orcid, publication_key(pub))] = pub.get("first_seen") or ""
    has_history = bool(groups)  # False on the very first run of this version
    return saved, first_seen, has_history


def mark_first_seen(works, orcid, first_seen, has_history, today):
    for pub in works:
        known = first_seen.get((orcid, publication_key(pub)))
        if known is not None:
            pub["first_seen"] = known
        else:
            # On the first run everything is already "old", so nothing gets a badge.
            pub["first_seen"] = today if has_history else ""


def main():
    OUTPUT_DIR.mkdir(exist_ok=True)
    roster = read_roster()
    previous, first_seen, has_history = load_previous()
    today = datetime.now(timezone.utc).date().isoformat()

    groups, failures, cache = {}, [], {}
    for group in GROUPS:
        info = roster[group["key"]]
        print(f"\n== {group['title']} ({len(info['people'])} with an ORCID iD) ==")
        for note in info["skipped"]:
            print(f"  skipped - {note}")

        researchers, used_slugs = [], set()
        for person in info["people"]:
            orcid = person["orcid"]
            status = "ok"
            if orcid in cache:
                works = cache[orcid]
            else:
                try:
                    works = get_orcid_works(orcid)
                    mark_first_seen(works, orcid, first_seen, has_history, today)
                    cache[orcid] = works
                except Exception as error:
                    works = previous.get(orcid, [])
                    status = "stale" if works else "error"
                    failures.append({"name": person["name"], "orcid": orcid, "error": str(error)})
                    print(f"  WARNING {person['name']}: {error}")
                time.sleep(REQUEST_PAUSE_SECONDS)

            slug = slugify(person["name"])
            while slug in used_slugs:
                slug += "-x"
            used_slugs.add(slug)

            print(f"  {person['name']:<32} {orcid}  {len(works):>4} works")
            researchers.append({**person, "slug": slug, "status": status,
                                "work_count": len(works), "works": works})

        groups[group["key"]] = {"title": group["title"], "file": group["file"],
                                "skipped": info["skipped"], "researchers": researchers}

    total_people = sum(len(g["researchers"]) for g in groups.values())
    if total_people and len(failures) >= total_people:
        raise RuntimeError("Every ORCID request failed; leaving the existing pages unchanged.")

    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    JSON_OUTPUT.write_text(json.dumps({
        "generated_at": generated_at,
        "source": "ORCID public API v3.0 (no other sources)",
        "failures": failures,
        "groups": groups,
    }, indent=2, ensure_ascii=False), encoding="utf-8")

    for key, group in groups.items():
        (OUTPUT_DIR / group["file"]).write_text(render_page(key, groups, generated_at), encoding="utf-8")
        print(f"Created docs/{group['file']}")
    (OUTPUT_DIR / "index.html").write_text(render_redirect(GROUPS[0]["file"]), encoding="utf-8")
    # Tell GitHub Pages to serve these files as-is instead of running Jekyll on them.
    (OUTPUT_DIR / ".nojekyll").write_text("", encoding="utf-8")
    print(f"Created {JSON_OUTPUT.relative_to(ROOT)}")
    if failures:
        print(f"Warnings recorded: {len(failures)}")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"Build failed: {error}", file=sys.stderr)
        sys.exit(1)
